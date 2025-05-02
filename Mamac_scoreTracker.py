import tkinter as tk
from tkinter import messagebox
from openpyxl import Workbook, load_workbook
import os

window = tk.Tk()
window.title("Students Grade")
window.geometry("500x500")
window.configure(bg='#f5f5f5')



E_FILE = "gradesofstudents.xlsx"



if os.path.exists(E_FILE):
    grade_book = load_workbook(E_FILE)
    sheet = grade_book.active
else:
    grade_book = Workbook()
    sheet = grade_book.active
    sheet.append(["Name", "Grades", "Remarks"])
    grade_book.save(E_FILE)
    
input_frame = tk.LabelFrame(window, text="Enter Student Information", padx=10, pady=10, bg='#e6f2ff', fg='#003366', font=('Arial', 10, 'bold'))
input_frame.pack(pady=10, fill="x")

tk.Label(input_frame, text="Name:",  bg='#B6D0E2', fg='#003366').grid(row=0, column=0, sticky="e")
name_entry = tk.Entry(input_frame, width=30)
name_entry.grid(row=0, column=1)

tk.Label(input_frame, text="Grades:",  bg='#B6D0E2', fg='#003366').grid(row=1, column=0, sticky="e")
score_entry = tk.Entry(input_frame, width=30)
score_entry.grid(row=1, column=1)

list_frame = tk.LabelFrame(window, text="Student Score Tracker", padx=10, pady=10, bg='#e6f2ff', fg='#003366', font=('Arial', 10, 'bold'))
list_frame.pack(pady=10, fill="both", expand=True)

student_list = tk.Listbox(list_frame, width=60, height=15, bg='white', fg='black',
                         selectbackground='#cce6ff', selectforeground='black')
student_list.pack(fill="both", expand=True)


def save_student():
    name = name_entry.get().strip()
    grades = score_entry.get().strip()
    
    if not name or not grades:
        messagebox.showerror("Error", "Please complete both inputs.")
        return
    
    try:
        grades = float(grades)
        if grades < 0 or grades > 100:
            messagebox.showerror("Error", "Grade must be between 0-100")
            return
    except ValueError:
        messagebox.showerror("Error", "Grade must be a number")
        return
    
    remarks = "Passed" if grades >= 75 else "Failed"
    
    sheet.append([name, grades, remarks])
    grade_book.save(E_FILE)
    
    student_list.insert(tk.END, f"{name:<20} {grades:<5}")
        
    name_entry.delete(0, tk.END)
    score_entry.delete(0, tk.END)
    name_entry.focus()
    
button_frame = tk.Frame(window)
button_frame.pack(pady=10)

save_button = tk.Button(button_frame, text="✓ Save!", command=save_student, bg='lightblue', fg='black',
                       activeforeground='white', relief='raised', padx=10)
save_button.pack(side="left", padx=5)





window.mainloop()